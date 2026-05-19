import argparse
from os import walk
from os.path import join

from sys import exc_info


from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from Components.Immunogenic_Epitopes.ImmunogenicEpitopesValidator import get_column_names


def parse_arguments():
    parser = argparse.ArgumentParser()
    parser.add_argument("-i", "--input", required=True, help="Input Directories, comma separated", type=str)

    return parser.parse_args()


def parse_input_directories(input_directory_string):
    input_directories = input_directory_string.split(',')
    file_list = []
    for input_directory in input_directories:
        for root, dirs, files in walk(input_directory):
            for file in files:
                file_list.append(join(root, file))
    file_list = [file for file in file_list if '.DS_Store' not in file]
    return file_list


def sort_files(upload_file_names):
    sorted_files = {'HML':[]
        ,'HAML':[]
        ,'CONTENT_SUMMARY':[]
        ,'ANTIBODY_CSV':[]
        ,'VALIDATION_REPORT':[]
        ,'PROJECT_DATA_MATRIX':[]
        ,'OTHER':[]}

    for upload_file_name in upload_file_names:
        if upload_file_name.endswith('.xml') and '_HML_' in upload_file_name:
            sorted_files['HML'].append(upload_file_name)
        elif upload_file_name.endswith('.hml') and '_HML_' in upload_file_name:
            sorted_files['HML'].append(upload_file_name)
        elif 'ContentSummary.csv' in upload_file_name:
            sorted_files['CONTENT_SUMMARY'].append(upload_file_name)
        elif upload_file_name.endswith('.haml') and '_HAML_' in upload_file_name:
            sorted_files['HAML'].append(upload_file_name)
        elif upload_file_name.endswith('.csv') and '_ANTIBODY_CSV_' in upload_file_name:
            sorted_files['ANTIBODY_CSV'].append(upload_file_name)
        elif upload_file_name.endswith('.csv') and '_INFO_CSV_' in upload_file_name:
            # I Guess some were uploaded as "INFO_CSV" files?
            sorted_files['ANTIBODY_CSV'].append(upload_file_name)
        elif upload_file_name.endswith('.csv') and '_HAML_' in upload_file_name:
            # Special case, some folks uploaded the .csv as a HAML data-type, oops.
            sorted_files['ANTIBODY_CSV'].append(upload_file_name)
        elif upload_file_name.endswith('.csv') and '_null_' in upload_file_name:
            # Some were uploaded without a file type, somehow?
            sorted_files['ANTIBODY_CSV'].append(upload_file_name)
        elif upload_file_name.endswith('.Validation_Report.xlsx') and '_PROJECT_DATA_MATRIX_' in upload_file_name:
            sorted_files['VALIDATION_REPORT'].append(upload_file_name)
        elif upload_file_name.endswith('.xlsx') and '_PROJECT_DATA_MATRIX_' in upload_file_name:
            sorted_files['PROJECT_DATA_MATRIX'].append(upload_file_name)
        elif upload_file_name.endswith('.xlsx') and '_PED_' in upload_file_name:
            # I don't know what to do with these.
            sorted_files['OTHER'].append(upload_file_name)
        else:
            raise Exception(f'Unknown file type:{upload_file_name}')
    return sorted_files


def read_content_summary(sorted_files):
    if len(sorted_files['CONTENT_SUMMARY']) != 1:
        raise Exception (f'There should be exactly one content summary file, instead there are {sorted_files["CONTENT_SUMMARY"]}')
    content_summary_filename = sorted_files['CONTENT_SUMMARY'][0]

    content_summary = []

    with open(content_summary_filename,'r') as content_summary_file:
        column_headers = None
        for line_index, line in enumerate(content_summary_file.readlines()):
            if line_index in [0,1]:
                print(f'Header Line:{line}')
            elif line_index == 2:
                column_headers = line.strip().split(',')
            else:
                try:
                    column_tokens = line.strip().split(',')
                    upload_data = {}

                    #Oops, extra commas in the data. Re-join them.
                    if (column_tokens[0] == '728: Faculty of Medicine Ramathibodi Hospital'
                    and column_tokens[1] == ' Mahidol University'):
                        column_tokens[0] = '728: Faculty of Medicine Ramathibodi Hospital; Mahidol University'
                        column_tokens.pop(1)

                    for i in range(0,len(column_tokens)):
                        upload_data[column_headers[i]] = column_tokens[i]
                    content_summary.append(upload_data)
                except Exception as e:
                    print(f'Something went wrong parsing content summary line:\n{line}')

    return content_summary


def read_data_matrix(data_matrix_filename):
    print(f'Reading data matrix {data_matrix_filename}')
    epitope_column_names = get_column_names(is_immunogenic=True)
    xl_workbook = load_workbook(filename=data_matrix_filename)

    input_columns = {}
    first_sheet = xl_workbook[xl_workbook.sheetnames[0]]
    for col_index_raw, col in enumerate(first_sheet.iter_cols()):
        if(col[0].value is not None and len(str(col[0].value))>0):
            header_name = str(col[0].value).strip().lower()
            column_letter = get_column_letter(col[0].column)
            input_columns[header_name] = column_letter

    # I don't know if i need these, yet
    missing_columns = [column for column in epitope_column_names if column not in input_columns]
    extra_columns = [column for column in input_columns if column not in epitope_column_names]

    return xl_workbook, missing_columns, extra_columns


def read_data_matrices(sorted_files):
    data_matrix_filenames = sorted_files['PROJECT_DATA_MATRIX']
    data_matrices = {}

    for data_matrix_filename in data_matrix_filenames:
        xl_workbook, missing_columns, extra_columns = read_data_matrix(data_matrix_filename=data_matrix_filename)
        data_matrices[data_matrix_filename] = xl_workbook

    return data_matrices


def connect_hla_genotypes(sorted_files, content_summary, data_matrices):
    for data_matrix_filename in data_matrices:
        print(f'Lookup HLA in {data_matrix_filename}')



def main(args):
    upload_file_names = parse_input_directories(input_directory_string=args.input)
    
    sorted_files = sort_files(upload_file_names=upload_file_names)

    content_summary = read_content_summary(sorted_files=sorted_files)

    data_matrices = read_data_matrices(sorted_files=sorted_files)

    data_with_connected_genotypes = connect_hla_genotypes(sorted_files=sorted_files
                                                          ,content_summary=content_summary
                                                          ,data_matrices=data_matrices)

    print('done')

if __name__=='__main__':
    try:
        args=parse_arguments()
        main(args=args)
    except Exception as e:
        print(f'Exception:{e}')
        print (str(exc_info()))
        raise e